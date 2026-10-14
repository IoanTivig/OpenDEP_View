import html
import io

from matplotlib.figure import Figure
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.func.general import format_frequency
from src.func.group_statistics import stars

'''
The statistics report of the replicate groups (group_statistics.build_report) as files for general use:
    HTML for the PDF (laid out and printed with Qt's QTextDocument / QPdfWriter, see ui/helpers/report_pdf.py)
    Excel workbook (.xlsx; opens in Excel, LibreOffice Calc and other office suites)
'''

TITLE = "OpenDEP View: statistics report of the replicate groups"

METHODS = [
    "Unit of observation: the biological replicate (suspension); its technical replicates are averaged first, "
    "so n is the number of suspensions. A group with a single suspension uses its technical replicates "
    "(level \"technical\": measurement spread, not biological spread).",
    "Values compared: the first and second crossover and the model parameters of the curve fitted to each "
    "replicate's scatter. Parameters are compared only between groups that use the same model. Values that are the "
    "same in every sample (e.g. fixed in the fits) are listed but not tested.",
    "Scale: crossovers and fitted parameters of repeated runs are log-normal (multiplicative spread), so every "
    "quantity whose values are all > 0 is analysed on the log scale and summarised as geometric mean ×/÷ geometric SD "
    "(about 68 % of the samples lie between mean / GSD and mean × GSD), with the 95 % confidence interval of the "
    "geometric mean (Limpert & Stahel, PLoS One 6 (2011) e21403; Bland & Altman, BMJ 312 (1996) 770). A quantity "
    "with a value ≤ 0 stays on the linear scale (mean ± SD), and so do the conductances that can be zero (membrane, "
    "nuclear envelope, surface conductance), which the fits put at their lower bound when they are not determined. "
    "The scale is set per quantity for all groups before any test. The ordinary mean ± SD is listed too, for "
    "comparison with published values.",
    "Normality: Shapiro-Wilk test per group (n ≥ 3) on the analysed scale (log values on the log scale); normal when "
    "p > α.",
    "Pairwise tests (both reported): Welch's t-test (normal data, unequal variances allowed; on log values it "
    "compares the geometric means, Bland & Altman, BMJ 312 (1996) 1153) and the Mann-Whitney U test (any "
    "distribution, two-sided, exact for small samples; ranks, so the same on both scales). The test used is Welch's "
    "t-test unless normality is rejected in either group, then Mann-Whitney U. With n < 3 normality cannot be tested "
    "and Welch's t-test is used.",
    "Multiple comparisons: the p values of the test used are corrected over all pairs of groups with the Holm "
    "method, separately for each quantity. Significance: * p < 0.05, ** p < 0.01, *** p < 0.001 (Holm).",
    "Effect size: the ratio of the geometric means (first group / second) with its 95 % confidence interval (Welch), "
    "or the difference of the means on the linear scale; and Hedges' g, the bias-corrected difference of the means "
    "in pooled SDs on the analysed scale (sign: first group minus second).",
    "Ranking of the pairs: as chosen in Best separation (Rank by), by the separation score (equal scores: larger "
    "ΔRe[CM] first), by the crossover ratio, or by ΔRe[CM] at the best frequency (pairs not on opposite sides there "
    "last); only the order changes, every pair's values are those at its own best frequency.",
    "Separation: the share of each group in pDEP at a frequency, from the first "
    "crossovers of its samples (each smoothed by a log-normal of width 0.2 ln). The score is the largest product of "
    "the difference of the two shares and the force weight min(1, ΔRe[CM] / 0.5), which is 0 unless the two groups' "
    "Re[CM] have opposite signs (0: no usable window, 1: one group entirely in pDEP, the other in nDEP, with a strong "
    "force). It is searched only inside the frequency range both groups were measured over, below their second "
    "crossovers; the separation at the second crossovers is listed apart when both groups have them inside that "
    "range. Samples on side: how many samples of each group are in pDEP (first group) and nDEP (second group) at the "
    "best frequency, from their crossovers; a few overlapping samples lower the score even when the crossover ratio "
    "is large. Crossover ratio: higher / lower geometric-mean first crossover. ΔRe[CM]: Re[CM] of the pDEP group "
    "minus the nDEP group at the best frequency, from the groups' curves (model with the geometric-mean parameters).",
    "Small samples: with 3 suspensions per group the Mann-Whitney U test cannot reach p < 0.1, and the "
    "Shapiro-Wilk test has little power; read the tests together with the effect sizes and the separation.",
]


def number(value, digits=4):
    return "–" if value is None else f"{value:.{digits}g}"


def p_text(p):
    return "–" if p is None else "< 0.001" if p < 0.001 else f"{p:.3f}"


def _frequency(value):
    return "–" if value is None else format_frequency(value)


# A quantity's mean with its spread: "31.15 ×/÷ 2.01" (log scale) or "0.84 ± 0.12" (linear scale)
def mean_text(q):
    if q["mean"] is None:
        return "–"
    if q["scale"] == "log":
        return f"{number(q['mean'])} ×/÷ {q['gsd']:.2f}" if q["gsd"] is not None else number(q["mean"])
    return f"{number(q['mean'])} ± {number(q['sd'], 2)}" if q["sd"] is not None else number(q["mean"])


def ordinary_text(q):
    if q["arithmetic_mean"] is None:
        return "–"
    sd = q["arithmetic_sd"]
    return number(q["arithmetic_mean"]) + (f" ± {number(sd, 2)}" if sd is not None else "")


def ci_text(q):
    return "–" if q["ci_low"] is None else f"{number(q['ci_low'])} – {number(q['ci_high'])}"


# Effect of a test: "2.45× (1.61–3.73)" (ratio of the geometric means) or "+0.12 (−0.03 to +0.27)" (difference)
def effect_text(t):
    if t["effect"] is None:
        return "–"
    if t["effect_kind"] == "ratio":
        interval = f" ({t['effect_low']:.3g}–{t['effect_high']:.3g})" if t["effect_low"] is not None else ""
        return f"{t['effect']:.3g}×{interval}"
    interval = f" ({t['effect_low']:+.3g} to {t['effect_high']:+.3g})" if t["effect_low"] is not None else ""
    return f"{t['effect']:+.3g}{interval}"


def side_text(pair):
    if not pair.get("a_side") or not pair.get("b_side"):
        return "–"
    return f"{pair['a_side'][0]}/{pair['a_side'][1]} · {pair['b_side'][0]}/{pair['b_side'][1]}"


def _model_line(report):
    return {entry["name"]: entry["model_name"] for entry in report["groups"]}


def significant_count(pair):
    tested = [t for t in pair["tests"].values() if t["p_holm"] is not None]
    return sum(1 for t in tested if t["p_holm"] < 0.05), len(tested)


# ---------------------------------------------------------------- figure
# The groups for the report (PNG bytes): each group's mean-parameter curve and mean ± SD points, and the best
# frequency of the best-separated pair. plots: [{"name", "color", "curve": (frequencies, recm) or None,
# "points": (frequencies, values, errors) or None}]
def report_figure(plots, best_frequency=None, best_label=""):
    figure = Figure(figsize=(9.0, 3.6), dpi=150)
    axes = figure.add_subplot()
    for plot in plots:
        if plot["curve"] is not None:
            axes.plot(plot["curve"][0], plot["curve"][1], color=plot["color"], linewidth=1.6, label=plot["name"])
        if plot["points"] is not None:
            frequencies, values, errors = plot["points"]
            axes.errorbar(frequencies, values, yerr=errors, fmt="o", markersize=3.5, color=plot["color"],
                          ecolor="#64748B", elinewidth=0.8, capsize=2,
                          label=None if plot["curve"] is not None else plot["name"])
    if best_frequency:
        axes.axvline(best_frequency, color="#0F172A", linestyle="--", linewidth=0.9, label=best_label)
    axes.axhline(0, color="#94A3B8", linewidth=0.6)
    axes.set_xscale("log")
    axes.set_xlabel("Frequency (Hz)")
    axes.set_ylabel("Re[CM]")
    axes.grid(True, which="major", color="#E2E8F0", linewidth=0.6)
    for side in ("top", "right"):
        axes.spines[side].set_visible(False)
    axes.legend(frameon=False, fontsize=7.5, loc="center left", bbox_to_anchor=(1.01, 0.5))
    figure.tight_layout()
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png")
    return buffer.getvalue()


# ---------------------------------------------------------------- HTML (PDF)
_CSS = """
body { font-family: 'Segoe UI', 'DejaVu Sans', 'Liberation Sans', sans-serif; font-size: 8pt; color: #0F172A; }
h1 { font-size: 15pt; font-weight: 600; margin-bottom: 2px; }
h2 { font-size: 11.5pt; font-weight: 600; margin-top: 14px; margin-bottom: 4px; }
h3 { font-size: 9.5pt; font-weight: 600; margin-top: 10px; margin-bottom: 3px; }
p.muted { color: #64748B; margin-top: 0px; }
p.note { color: #334155; margin-top: 2px; margin-bottom: 2px; }
table { border-collapse: collapse; }
th { background-color: #E2E8F0; font-weight: 600; text-align: left; }
td.num, th.num { text-align: right; }
td.sig { background-color: #DCFCE7; }
"""


def _table(headers, rows, numeric=()):
    head = "".join(f"<th class='num'>{h}</th>" if i in numeric else f"<th>{h}</th>" for i, h in enumerate(headers))
    body = []
    for row in rows:
        cells = []
        for i, cell in enumerate(row):
            text, cls = (cell if isinstance(cell, tuple) else (cell, ""))
            classes = " ".join(c for c in (("num" if i in numeric else ""), cls) if c)
            cells.append(f"<td class='{classes}'>{text}</td>" if classes else f"<td>{text}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return ("<table border='1' cellspacing='0' cellpadding='3' width='100%' style='border-color:#CBD5E1'>"
            f"<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>")


def _dot(color):
    return f"<span style='color:{color}'>●</span>"


# The report as HTML; image: the resource name of the graph (added to the QTextDocument), or None
def report_html(report, project_name="", image=None, image_width=820):
    e = html.escape
    groups = report["groups"]
    parts = [f"<html><head><style>{_CSS}</style></head><body>", f"<h1>{TITLE}</h1>",
             f"<p class='muted'>{e(project_name) + ' · ' if project_name else ''}{report['created']} · "
             f"{len(report['groups'])} groups · {len(report['pairs'])} pairs · α = {report['alpha']}</p>"]
    if image:                                         # first on the page: an image is not split over two pages
        parts.append(f"<p align='center'><img src='{image}' width='{image_width}'></p>")

    # 1. Groups
    parts.append("<h2>1. Groups</h2>")
    parts.append(_table(["Group", "Model", "Biological replicates", "Technical replicates"],
                        [[f"{_dot(g['color'])} {e(g['name'])}", e(g["model_name"]), g["n_biological"],
                          g["n_technical"]] for g in report["groups"]], numeric=(2, 3)))
    if report["fixed"]:
        fixed = ", ".join(f"{e(f['label'])} {number(f['value'])} {e(f['unit'])}" for f in report["fixed"])
        parts.append(f"<p class='note'>Same value in every sample (not tested): {fixed}</p>")

    # 2. Separation ranking
    parts.append(f"<h2>2. Separation of the pairs (ranked by {e(report['rank_label'])})</h2>")
    rows = []
    for pair in report["pairs"]:
        significant, tested = significant_count(pair)
        rows.append([pair["rank"] or "–", f"{_dot(groups[pair['a_index']]['color'])} {e(pair['a'])}",
                     f"{_dot(groups[pair['b_index']]['color'])} {e(pair['b'])}", _frequency(pair["frequency"]),
                     number(pair["score"], 2) if pair["score"] is not None else "–",
                     f"{pair['split']:.2f} × {pair['weight']:.2f}" if pair["weight"] is not None else
                     number(pair["split"], 2) if pair["split"] is not None else "–",
                     side_text(pair),
                     f"{pair['ratio']:.2f}×" if pair["ratio"] else "–",
                     f"{pair['contrast']:.2f} ({pair['recm_a']:+.2f} / {pair['recm_b']:+.2f})"
                     if pair["contrast"] is not None else "–",
                     f"{pair['second']['score']:.2f} at {_frequency(pair['second']['frequency'])}"
                     if pair["second"] else "–",
                     f"{significant} / {tested}" if tested else "–"])
    parts.append(_table(["Rank", "pDEP group", "nDEP group", "Best frequency", "Separation score",
                         "Share difference × force weight", "Samples on side (pDEP · nDEP)", "Crossover ratio",
                         "ΔRe[CM] (pDEP / nDEP)", "2nd crossover", "Significant differences"],
                        rows, numeric=(0, 3, 4, 5, 6, 7, 8, 9, 10)))
    parts.append("<p class='note'>Samples on side: samples of the pDEP group in pDEP and of the nDEP group in nDEP at "
                 "the best frequency. Significant differences: quantities with p &lt; 0.05 (Holm) of those tested. "
                 "2nd crossover: the separation at the high-frequency crossovers, when both groups have them inside "
                 "the measured range (not used for the ranking).</p>")

    # 3. Descriptive statistics
    parts.append("<h2>3. Descriptive statistics</h2>")
    parts.append("<p class='note'>Mean ×/÷ GSD: geometric mean and geometric SD (log scale; about 68 % of the samples "
                 "lie between mean / GSD and mean × GSD). Mean ± SD: quantities on the linear scale (a value ≤ 0, "
                 "or a conductance that can be zero). "
                 "95 % CI: of the (geometric) mean. Shapiro-Wilk on the analysed scale. Ordinary mean ± SD: for "
                 "comparison with published values.</p>")
    if report["fixed"]:
        parts.append("<p class='note'>Values that are the same in every sample are listed under 1. Groups.</p>")
    fixed = {f["key"] for f in report["fixed"]}
    for g in report["groups"]:
        parts.append(f"<h3>{_dot(g['color'])} {e(g['name'])} · {e(g['model_name'])}</h3>")
        rows = []
        for key, q in g["quantities"].items():
            if q["n"] == 0 or key in fixed:
                continue
            normal = {True: "yes", False: ("no", "sig"), None: "–"}[q["normal"]]
            rows.append([f"{e(q['label'])} ({e(q['unit'])})", f"{q['n']} {q['level'][:4]}.", e(mean_text(q)),
                         ci_text(q), number(q["median"]), f"{number(q['min'])} – {number(q['max'])}",
                         number(q["shapiro_w"], 3), e(p_text(q["shapiro_p"])), normal, e(ordinary_text(q))])
        parts.append(_table(["Quantity", "n", "Mean ×/÷ GSD (or ± SD)", "95 % CI", "Median", "Range", "Shapiro W",
                             "Shapiro p", "Normal", "Ordinary mean ± SD"], rows, numeric=(2, 3, 4, 6, 7, 9)))

    # 4. Pairwise tests, in the order of the separation ranking
    parts.append("<h2 style='page-break-before: always'>4. Pairwise tests (pairs in the order of the separation "
                 "ranking)</h2>")
    for pair in report["pairs"]:
        rank = f"{pair['rank']}. " if pair["rank"] else ""
        separation = (f" · score {pair['score']:.2f} at {_frequency(pair['frequency'])}"
                      if pair["score"] is not None else " · no separation (crossovers missing)")
        a, b = groups[pair["a_index"]], groups[pair["b_index"]]
        parts.append(f"<h3>{rank}{_dot(a['color'])} {e(pair['a'])} vs {_dot(b['color'])} "
                     f"{e(pair['b'])}{separation}</h3>")
        if not pair["same_model"]:
            parts.append("<p class='note'>Different models: only the crossovers are compared.</p>")
        rows = []
        for key, t in pair["tests"].items():
            qa, qb = a["quantities"][key], b["quantities"][key]
            significant = t["p_holm"] is not None and t["p_holm"] < 0.05
            rows.append([f"{e(t['label'])} ({e(t['unit'])})", e(mean_text(qa)), e(mean_text(qb)),
                         number(t["welch_t"], 3), e(p_text(t["welch_p"])), number(t["mwu_u"], 3),
                         e(p_text(t["mwu_p"])), t["used"] or e(t["reason"]),
                         (e(p_text(t["p_holm"])), "sig" if significant else ""),
                         stars(t["p_holm"]) if t["p_holm"] is not None else "", e(effect_text(t)),
                         number(t["g"], 3)])
        if rows:
            parts.append(_table([f"Quantity", f"{e(pair['a'])}", f"{e(pair['b'])}", "Welch t", "Welch p", "MW U",
                                 "MW p", "Test used", "p (Holm)", "", "Ratio A / B, or A − B (95 % CI)",
                                 "Hedges g"],
                                rows, numeric=(1, 2, 3, 4, 5, 6, 8, 10, 11)))
        else:
            parts.append("<p class='note'>Nothing to test.</p>")

    # 5. Methods
    parts.append("<h2>5. Methods</h2>")
    parts.extend(f"<p class='note'>{e(line)}</p>" for line in METHODS)
    parts.append("</body></html>")
    return "".join(parts)


# ---------------------------------------------------------------- Excel workbook
_HEADER_FILL = PatternFill("solid", start_color="E2E8F0")
_SIGNIFICANT_FILL = PatternFill("solid", start_color="DCFCE7")
_THIN = Side(style="thin", color="CBD5E1")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _sheet_table(sheet, top, headers, rows, widths=None):
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=top, column=col, value=header)
        cell.font, cell.fill, cell.border = Font(bold=True), _HEADER_FILL, _BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, row in enumerate(rows, top + 1):
        for col, value in enumerate(row, 1):
            fill = None
            if isinstance(value, tuple):
                value, fill = value
            cell = sheet.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            if fill:
                cell.fill = fill
            if isinstance(value, float):
                cell.number_format = "0.000E+00" if value and (abs(value) < 1e-3 or abs(value) >= 1e5) else "0.0000"
    for col, width in enumerate(widths or [], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    return top + len(rows) + 1


def write_xlsx(report, path, project_name=""):
    book = Workbook()

    # Summary: groups, fixed values, methods
    sheet = book.active
    sheet.title = "Summary"
    sheet["A1"] = TITLE
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = (f"{project_name + ' · ' if project_name else ''}{report['created']} · α = {report['alpha']} · "
                   f"pairs ranked by {report['rank_label']}")
    row = _sheet_table(sheet, 4, ["Group", "Model", "Biological replicates", "Technical replicates"],
                       [[g["name"], g["model_name"], g["n_biological"], g["n_technical"]] for g in report["groups"]],
                       widths=[28, 24, 20, 20])
    if report["fixed"]:
        row = _sheet_table(sheet, row + 1, ["Same value in every sample (not tested)", "Value", "Unit"],
                           [[f["label"], f["value"], f["unit"]] for f in report["fixed"]])
    sheet.cell(row=row + 1, column=1, value="Methods").font = Font(bold=True)
    for i, line in enumerate(METHODS, row + 2):
        sheet.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=False)

    # Separation ranking
    sheet = book.create_sheet("Separation")
    rows = []
    for pair in report["pairs"]:
        significant, tested = significant_count(pair)
        second = pair["second"] or {}
        a_side, b_side = pair.get("a_side") or (None, None), pair.get("b_side") or (None, None)
        rows.append([pair["rank"], pair["a"], pair["b"], pair["frequency"], pair["score"], pair["split"],
                     pair["weight"], a_side[0], a_side[1], b_side[0], b_side[1], pair["ratio"], pair["contrast"],
                     pair["recm_a"], pair["recm_b"], second.get("score"), second.get("frequency"),
                     second.get("pdep"), significant if tested else None, tested or None])
    _sheet_table(sheet, 1, ["Rank", "pDEP group", "nDEP group", "Best frequency (Hz)", "Separation score",
                            "pDEP share difference", "Force weight", "pDEP group: samples in pDEP", "pDEP group: n",
                            "nDEP group: samples in nDEP", "nDEP group: n", "Crossover ratio", "ΔRe[CM]",
                            "Re[CM] pDEP group", "Re[CM] nDEP group", "2nd crossover score",
                            "2nd crossover frequency (Hz)", "2nd crossover pDEP group",
                            "Significant (p Holm < 0.05)", "Tested quantities"], rows,
                 widths=[7, 22, 22, 18, 16, 16, 13, 16, 9, 16, 9, 15, 12, 18, 18, 18, 22, 22, 18, 16])
    sheet.freeze_panes = "B2"

    # Descriptive statistics
    sheet = book.create_sheet("Descriptive")
    rows = []
    for g in report["groups"]:
        for q in g["quantities"].values():
            if q["n"]:
                rows.append([g["name"], q["label"], q["unit"], q["level"], q["n"], q["scale"], q["mean"], q["gsd"],
                             q["sd"], q["ci_low"], q["ci_high"], q["median"], q["min"], q["max"], q["shapiro_w"],
                             q["shapiro_p"],
                             {True: "yes", False: ("no", _SIGNIFICANT_FILL), None: "not testable"}[q["normal"]],
                             q["arithmetic_mean"], q["arithmetic_sd"]])
    _sheet_table(sheet, 1, ["Group", "Quantity", "Unit", "Level", "n", "Scale", "Mean (geometric on the log scale)",
                            "GSD (×/÷, log scale)", "SD (linear scale)", "95 % CI low", "95 % CI high", "Median",
                            "Min", "Max", "Shapiro W", "Shapiro p", "Normal", "Ordinary mean", "Ordinary SD"], rows,
                 widths=[20, 26, 9, 11, 5, 8, 16, 12, 12, 13, 13, 13, 13, 13, 11, 11, 13, 14, 13])
    sheet.freeze_panes = "C2"

    # Pairwise tests
    sheet = book.create_sheet("Pairwise tests")
    rows = []
    for pair in report["pairs"]:
        a, b = report["groups"][pair["a_index"]], report["groups"][pair["b_index"]]
        for key, t in pair["tests"].items():
            significant = t["p_holm"] is not None and t["p_holm"] < 0.05
            qa, qb = a["quantities"][key], b["quantities"][key]
            spread = (lambda q: q["gsd"] if q["scale"] == "log" else q["sd"])
            rows.append([pair["rank"], pair["a"], pair["b"], t["label"], t["unit"], qa["scale"], qa["mean"],
                         spread(qa), qb["mean"], spread(qb), t["welch_t"], t["welch_p"], t["mwu_u"], t["mwu_p"],
                         t["used"] or t["reason"],
                         (t["p_holm"], _SIGNIFICANT_FILL) if significant else t["p_holm"], stars(t["p_holm"]),
                         t["effect_kind"], t["effect"], t["effect_low"], t["effect_high"], t["g"]])
    _sheet_table(sheet, 1, ["Rank", "Group A (pDEP)", "Group B (nDEP)", "Quantity", "Unit", "Scale", "Mean A",
                            "GSD / SD A", "Mean B", "GSD / SD B", "Welch t", "Welch p", "Mann-Whitney U",
                            "Mann-Whitney p", "Test used", "p (Holm)", "Significance", "Effect", "Ratio A/B or A − B",
                            "95 % CI low", "95 % CI high", "Hedges g"], rows,
                 widths=[7, 18, 18, 26, 9, 8, 13, 11, 13, 11, 11, 11, 15, 15, 34, 11, 12, 10, 16, 12, 12, 11])
    sheet.freeze_panes = "E2"

    # Sample values (one per biological replicate, or per technical replicate in a single-suspension group)
    sheet = book.create_sheet("Sample values")
    rows = []
    for g in report["groups"]:
        for key, values in g["values"].items():
            q = g["quantities"][key]
            if values:
                rows.append([g["name"], q["label"], q["unit"], q["level"], *values])
    width = max((len(r) for r in rows), default=4)
    _sheet_table(sheet, 1, ["Group", "Quantity", "Unit", "Level"] + [f"Sample {i}" for i in range(1, width - 3)],
                 rows, widths=[20, 26, 9, 11] + [13] * (width - 4))
    book.save(path)
